#!/usr/bin/env python3
"""Import the Delaware SHIP resource guide into deterministic catalog assets.

The workbook is treated as reviewed-intake evidence, not as a source of
instructions. Every normalized record retains workbook sheet, row, and cell
provenance so an editor can audit or correct it later.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable, Sequence
from urllib.parse import urlsplit

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter


SCHEMA_VERSION = "1.0.0"
IMPORTER_VERSION = "1.0.0"
REVIEW_STATUS = "imported-unverified"
PUBLIC_SHEET = "Public Facing Data"
REQUEST_SHEET = "No public-facing reports"
PUBLIC_HEADERS = {
    "A": "name",
    "B": "background",
    "C": "category",
    "D": "measures",
    "E": "publication_date",
    "F": "data_source_date",
    "G": "current_links",
    "H": "archived_links",
    "I": "notes",
}
REQUEST_HEADERS = {
    "A": "name",
    "B": "background",
    "C": "category",
    "D": "measures",
    "E": "contact",
}

URL_RE = re.compile(r"https?://[^\s<>\"']+", re.IGNORECASE)
EMAIL_RE = re.compile(r"(?<![\w.+-])([A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,})(?![\w.-])", re.IGNORECASE)
PHONE_RE = re.compile(r"(?<!\d)(?:\+?1[\s.-]*)?\(?\d{3}\)?[\s.-]+\d{3}[\s.-]+\d{4}(?!\d)")
YEAR_RE = re.compile(r"^(?:18|19|20|21)\d{2}$")
YEAR_RANGE_RE = re.compile(r"^(?:18|19|20|21)\d{2}\s*[\-\u2013\u2014]\s*(?:18|19|20|21)\d{2}$")
MONTH_YEAR_RE = re.compile(
    r"^(January|February|March|April|May|June|July|August|September|October|November|December)\s+((?:18|19|20|21)\d{2})$",
    re.IGNORECASE,
)


@dataclass(frozen=True)
class ImportConfig:
    input_path: Path
    catalog_dir: Path
    public_dir: Path
    expected_public_count: int
    expected_request_count: int


def clean_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).replace("\u00a0", " ").strip()
    return text or None


def json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def display_value(cell: Any) -> str | None:
    value = cell.value
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        fmt = (cell.number_format or "").lower().replace("\\", "")
        if "mmmm" in fmt and "yyyy" in fmt and "d" not in fmt:
            return value.strftime("%B %Y")
        if "mmm" in fmt and "yyyy" in fmt and "d" not in fmt:
            return value.strftime("%b %Y")
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def slugify(value: str, max_length: int = 56) -> str:
    normalized = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    slug = re.sub(r"[^a-z0-9]+", "-", normalized.lower()).strip("-")
    return (slug or "item")[:max_length].rstrip("-")


def stable_id(prefix: str, *parts: Any, label: str | None = None) -> str:
    canonical = "\x1f".join("" if part is None else str(part).strip().casefold() for part in parts)
    digest = hashlib.sha256(canonical.encode("utf-8")).hexdigest()[:10]
    base = slugify(label or str(parts[0]) if parts else prefix)
    return f"{prefix}-{base}-{digest}"


def workbook_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def meaningful_last_row(sheet: Any, columns: Sequence[str]) -> int:
    for row in range(sheet.max_row, 0, -1):
        if any(clean_text(sheet[f"{column}{row}"].value) for column in columns):
            return row
    return 0


def merged_range_map(sheet: Any) -> dict[str, str]:
    result: dict[str, str] = {}
    for merged in sheet.merged_cells.ranges:
        for row in sheet.iter_rows(
            min_row=merged.min_row,
            max_row=merged.max_row,
            min_col=merged.min_col,
            max_col=merged.max_col,
        ):
            for cell in row:
                result[cell.coordinate] = str(merged)
    return result


def raw_cell_record(cell: Any, header: str, merged_ranges: dict[str, str]) -> dict[str, Any]:
    hyperlink = cell.hyperlink.target if cell.hyperlink else None
    return {
        "cell": cell.coordinate,
        "row": cell.row,
        "column": get_column_letter(cell.column),
        "field": header,
        "raw_value": json_value(cell.value),
        "display_value": display_value(cell),
        "data_type": cell.data_type,
        "number_format": cell.number_format,
        "hyperlink_target": hyperlink,
        "merged_range": merged_ranges.get(cell.coordinate),
    }


def collect_raw_cells(
    sheet: Any,
    row_start: int,
    row_end: int,
    headers: dict[str, str],
    merged_ranges: dict[str, str],
) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for row in range(row_start, row_end + 1):
        for column, header in headers.items():
            cell = sheet[f"{column}{row}"]
            if isinstance(cell, MergedCell) or cell.value is None:
                continue
            records.append(raw_cell_record(cell, header, merged_ranges))
    return records


def field_cells(sheet: Any, row_start: int, row_end: int, column: str) -> list[str]:
    cells: list[str] = []
    for row in range(row_start, row_end + 1):
        cell = sheet[f"{column}{row}"]
        if not isinstance(cell, MergedCell) and cell.value is not None:
            cells.append(cell.coordinate)
    return cells


def split_measures(value: Any) -> list[str]:
    text = clean_text(value)
    if not text:
        return []
    candidates = re.split(r"[;,\n]+", text)
    result: list[str] = []
    seen: set[str] = set()
    for candidate in candidates:
        item = re.sub(r"\s+", " ", candidate).strip(" \t\r\n.-")
        if not item:
            continue
        key = item.casefold()
        if key not in seen:
            seen.add(key)
            result.append(item)
    return result


def normalize_date_cell(cell: Any) -> dict[str, Any] | None:
    if cell.value is None:
        return None
    raw = json_value(cell.value)
    shown = display_value(cell)
    normalized: str | None = None
    precision = "unknown"
    if isinstance(cell.value, (datetime, date)):
        fmt = (cell.number_format or "").lower().replace("\\", "")
        if "yyyy" in fmt and "d" not in fmt:
            normalized = cell.value.strftime("%Y-%m") if "m" in fmt else cell.value.strftime("%Y")
            precision = "month" if "m" in fmt else "year"
        else:
            normalized = cell.value.date().isoformat() if isinstance(cell.value, datetime) else cell.value.isoformat()
            precision = "day"
    else:
        text = clean_text(cell.value) or ""
        if isinstance(cell.value, (int, float)) and float(cell.value).is_integer() and YEAR_RE.match(str(int(cell.value))):
            normalized = str(int(cell.value))
            precision = "year"
        elif YEAR_RE.match(text):
            normalized = text
            precision = "year"
        elif YEAR_RANGE_RE.match(text):
            normalized = re.sub(r"\s*[\u2013\u2014]\s*", "-", text)
            precision = "range"
        else:
            month_match = MONTH_YEAR_RE.match(text)
            if month_match:
                month = datetime.strptime(month_match.group(1), "%B").month
                normalized = f"{month_match.group(2)}-{month:02d}"
                precision = "month"
            elif text.casefold() in {"n.d.", "n.d", "unknown", "not available", "na", "n/a"}:
                normalized = None
                precision = "unknown"
            else:
                normalized = text
                precision = "label"
    return {
        "raw_value": raw,
        "display_value": shown,
        "normalized": normalized,
        "precision": precision,
        "cell": cell.coordinate,
    }


def unique_dicts(items: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    seen: set[str] = set()
    for item in items:
        key = json.dumps(item, ensure_ascii=False, sort_keys=True)
        if key not in seen:
            seen.add(key)
            result.append(item)
    return result


def extract_urls(text: Any) -> list[str]:
    value = clean_text(text)
    if not value:
        return []
    urls: list[str] = []
    seen: set[str] = set()
    for match in URL_RE.findall(value):
        url = match.rstrip(".,;:)]}")
        if url and url not in seen:
            seen.add(url)
            urls.append(url)
    return urls


def valid_http_url(url: str) -> bool:
    try:
        parsed = urlsplit(url)
        return parsed.scheme in {"http", "https"} and bool(parsed.netloc) and " " not in url
    except ValueError:
        return False


def infer_format(url: str) -> str:
    path = urlsplit(url).path.lower()
    extensions = {
        ".csv": "csv",
        ".json": "json",
        ".geojson": "geojson",
        ".xlsx": "xlsx",
        ".xls": "xls",
        ".pdf": "pdf",
        ".zip": "zip",
        ".xml": "xml",
        ".txt": "txt",
    }
    for extension, label in extensions.items():
        if path.endswith(extension):
            return label
    return "web"


def distribution_records(
    source_id: str,
    sheet: Any,
    row_start: int,
    row_end: int,
    column: str,
    role: str,
    topic_by_row: dict[int, str],
) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    merged_coordinates = {
        cell.coordinate
        for merged in sheet.merged_cells.ranges
        if merged.max_row > merged.min_row or merged.max_col > merged.min_col
        for row_cells in sheet.iter_rows(
            min_row=merged.min_row,
            max_row=merged.max_row,
            min_col=merged.min_col,
            max_col=merged.max_col,
        )
        for cell in row_cells
    }
    for row in range(row_start, row_end + 1):
        cell = sheet[f"{column}{row}"]
        if isinstance(cell, MergedCell) or cell.value is None:
            continue
        shown = display_value(cell)
        visible_urls = extract_urls(cell.value)
        hyperlink_target = cell.hyperlink.target if cell.hyperlink else None
        candidates = list(visible_urls)
        if not candidates and hyperlink_target and hyperlink_target.lower().startswith(("http://", "https://")):
            candidates.append(hyperlink_target)
        topic_id = None if cell.coordinate in merged_coordinates else topic_by_row.get(row)
        for url in candidates:
            records.append(
                {
                    "id": stable_id("dist", source_id, role, url, label=role),
                    "url": url,
                    "role": role,
                    "format": infer_format(url),
                    "topic_id": topic_id,
                    "display_text": shown,
                    "hyperlink_target": hyperlink_target,
                    "provenance": {"sheet": sheet.title, "row": row, "cell": cell.coordinate},
                }
            )
    return records


def contact_record(cell: Any) -> dict[str, Any]:
    raw = display_value(cell) or ""
    hyperlink = cell.hyperlink.target if cell.hyperlink else None
    emails: list[str] = []
    for email in EMAIL_RE.findall(raw):
        if email.casefold() not in {item.casefold() for item in emails}:
            emails.append(email)
    if hyperlink and hyperlink.lower().startswith("mailto:"):
        linked_email = hyperlink[7:].split("?", 1)[0].split("%0", 1)[0]
        if EMAIL_RE.fullmatch(linked_email) and linked_email.casefold() not in {item.casefold() for item in emails}:
            emails.append(linked_email)
    phones: list[str] = []
    for phone in PHONE_RE.findall(raw):
        normalized = re.sub(r"\s+", " ", phone).strip()
        if normalized not in phones:
            phones.append(normalized)
    return {
        "raw": raw or None,
        "emails": emails,
        "phones": phones,
        "urls": extract_urls(raw),
        "hyperlink_target": hyperlink,
        "provenance": {"sheet": cell.parent.title, "row": cell.row, "cell": cell.coordinate},
    }


def make_topic(source_id: str, sheet: Any, row: int) -> dict[str, Any] | None:
    category = clean_text(sheet[f"C{row}"].value)
    measures_raw = clean_text(sheet[f"D{row}"].value)
    publication_date = normalize_date_cell(sheet[f"E{row}"]) if sheet.title == PUBLIC_SHEET else None
    data_source_date = normalize_date_cell(sheet[f"F{row}"]) if sheet.title == PUBLIC_SHEET else None
    if not any((category, measures_raw, publication_date, data_source_date)):
        return None
    topic_id = stable_id("topic", source_id, category, measures_raw, label=category or measures_raw or f"row-{row}")
    fields = {
        "category": f"C{row}" if sheet[f"C{row}"].value is not None else None,
        "measures": f"D{row}" if sheet[f"D{row}"].value is not None else None,
    }
    if sheet.title == PUBLIC_SHEET:
        fields.update(
            {
                "publication_date": f"E{row}" if sheet[f"E{row}"].value is not None else None,
                "data_source_date": f"F{row}" if sheet[f"F{row}"].value is not None else None,
            }
        )
    return {
        "id": topic_id,
        "category": category,
        "measures": split_measures(measures_raw),
        "measures_raw": measures_raw,
        "publication_date": publication_date,
        "data_source_date": data_source_date,
        "provenance": {"sheet": sheet.title, "row": row, "cells": fields},
    }


def public_group_ranges(sheet: Any) -> list[tuple[int, int]]:
    last_row = meaningful_last_row(sheet, tuple(PUBLIC_HEADERS))
    starts = [row for row in range(3, last_row + 1) if clean_text(sheet[f"A{row}"].value)]
    return [(start, starts[index + 1] - 1 if index + 1 < len(starts) else last_row) for index, start in enumerate(starts)]


def build_public_sources(sheet: Any, workbook_name: str) -> list[dict[str, Any]]:
    merged_ranges = merged_range_map(sheet)
    sources: list[dict[str, Any]] = []
    for row_start, row_end in public_group_ranges(sheet):
        name = clean_text(sheet[f"A{row_start}"].value)
        if not name:
            continue
        source_id = stable_id("src", name, "public", label=name)
        topics = [topic for row in range(row_start, row_end + 1) if (topic := make_topic(source_id, sheet, row))]
        topic_by_row = {topic["provenance"]["row"]: topic["id"] for topic in topics}
        distributions = distribution_records(source_id, sheet, row_start, row_end, "G", "current", topic_by_row)
        distributions.extend(distribution_records(source_id, sheet, row_start, row_end, "H", "archive", topic_by_row))
        notes = [clean_text(sheet[f"I{row}"].value) for row in range(row_start, row_end + 1)]
        notes = [note for note in notes if note]
        publication_dates = unique_dicts(
            topic["publication_date"] for topic in topics if topic["publication_date"] is not None
        )
        data_source_dates = unique_dicts(
            topic["data_source_date"] for topic in topics if topic["data_source_date"] is not None
        )
        sources.append(
            {
                "id": source_id,
                "name": name,
                "background": clean_text(sheet[f"B{row_start}"].value),
                "access_type": "public",
                "review_status": REVIEW_STATUS,
                "notes": notes,
                "topics": topics,
                "publication_dates": publication_dates,
                "data_source_dates": data_source_dates,
                "distributions": distributions,
                "contacts": [],
                "license": None,
                "monitoring": {"availability_status": "not_checked", "last_checked_at": None},
                "provenance": {
                    "source_workbook": workbook_name,
                    "sheet": sheet.title,
                    "row_start": row_start,
                    "row_end": row_end,
                    "fields": {
                        "name": field_cells(sheet, row_start, row_end, "A"),
                        "background": field_cells(sheet, row_start, row_end, "B"),
                        "category": field_cells(sheet, row_start, row_end, "C"),
                        "measures": field_cells(sheet, row_start, row_end, "D"),
                        "publication_date": field_cells(sheet, row_start, row_end, "E"),
                        "data_source_date": field_cells(sheet, row_start, row_end, "F"),
                        "current_links": field_cells(sheet, row_start, row_end, "G"),
                        "archived_links": field_cells(sheet, row_start, row_end, "H"),
                        "notes": field_cells(sheet, row_start, row_end, "I"),
                    },
                    "raw_cells": collect_raw_cells(
                        sheet, row_start, row_end, PUBLIC_HEADERS, merged_ranges
                    ),
                },
            }
        )
    return sources


def request_rows(sheet: Any) -> list[int]:
    last_row = meaningful_last_row(sheet, tuple(REQUEST_HEADERS))
    return [
        row
        for row in range(2, last_row + 1)
        if clean_text(sheet[f"A{row}"].value)
        and not (clean_text(sheet[f"A{row}"].value) or "").startswith("*")
    ]


def build_request_sources(sheet: Any, workbook_name: str) -> list[dict[str, Any]]:
    merged_ranges = merged_range_map(sheet)
    sources: list[dict[str, Any]] = []
    for row in request_rows(sheet):
        name = clean_text(sheet[f"A{row}"].value)
        if not name:
            continue
        source_id = stable_id("src", name, "request-only", label=name)
        topic = make_topic(source_id, sheet, row)
        topics = [topic] if topic else []
        contact = contact_record(sheet[f"E{row}"])
        distributions: list[dict[str, Any]] = []
        for url in contact["urls"]:
            distributions.append(
                {
                    "id": stable_id("dist", source_id, "request_info", url, label="request-info"),
                    "url": url,
                    "role": "request_info",
                    "format": infer_format(url),
                    "topic_id": topic["id"] if topic else None,
                    "display_text": contact["raw"],
                    "hyperlink_target": contact["hyperlink_target"],
                    "provenance": {"sheet": sheet.title, "row": row, "cell": f"E{row}"},
                }
            )
        sources.append(
            {
                "id": source_id,
                "name": name,
                "background": clean_text(sheet[f"B{row}"].value),
                "access_type": "request_only",
                "review_status": REVIEW_STATUS,
                "notes": [],
                "topics": topics,
                "publication_dates": [],
                "data_source_dates": [],
                "distributions": distributions,
                "contacts": [contact],
                "license": None,
                "monitoring": {"availability_status": "not_checked", "last_checked_at": None},
                "provenance": {
                    "source_workbook": workbook_name,
                    "sheet": sheet.title,
                    "row_start": row,
                    "row_end": row,
                    "fields": {
                        "name": field_cells(sheet, row, row, "A"),
                        "background": field_cells(sheet, row, row, "B"),
                        "category": field_cells(sheet, row, row, "C"),
                        "measures": field_cells(sheet, row, row, "D"),
                        "contact": field_cells(sheet, row, row, "E"),
                    },
                    "raw_cells": collect_raw_cells(sheet, row, row, REQUEST_HEADERS, merged_ranges),
                },
            }
        )
    return sources


def make_catalog(config: ImportConfig) -> dict[str, Any]:
    workbook = load_workbook(config.input_path, data_only=False, read_only=False)
    missing = [name for name in (PUBLIC_SHEET, REQUEST_SHEET) if name not in workbook.sheetnames]
    if missing:
        raise ValueError(f"Missing required worksheet(s): {', '.join(missing)}")
    public_sheet = workbook[PUBLIC_SHEET]
    request_sheet = workbook[REQUEST_SHEET]
    public_sources = build_public_sources(public_sheet, config.input_path.name)
    request_sources = build_request_sources(request_sheet, config.input_path.name)
    sources = public_sources + request_sources
    introduction = clean_text(public_sheet["B1"].value)
    request_note = clean_text(request_sheet[f"A{meaningful_last_row(request_sheet, tuple(REQUEST_HEADERS))}"].value)
    return {
        "schema_version": SCHEMA_VERSION,
        "importer_version": IMPORTER_VERSION,
        "catalog_type": "reviewed-intake",
        "review_status": REVIEW_STATUS,
        "source_workbook": {
            "filename": config.input_path.name,
            "sha256": workbook_sha256(config.input_path),
            "worksheets": {
                "public": PUBLIC_SHEET,
                "request_only": REQUEST_SHEET,
            },
            "context": introduction,
            "request_only_note": request_note,
        },
        "summary": {
            "source_count": len(sources),
            "public_source_count": len(public_sources),
            "request_only_source_count": len(request_sources),
            "topic_count": sum(len(source["topics"]) for source in sources),
            "distribution_count": sum(len(source["distributions"]) for source in sources),
        },
        "sources": sources,
    }


def duplicate_values(values: Iterable[str]) -> list[str]:
    counts = Counter(values)
    return sorted(value for value, count in counts.items() if count > 1)


def warning(issue: str, source: dict[str, Any], detail: str, cells: Sequence[str] | None = None) -> dict[str, Any]:
    return {
        "issue": issue,
        "source_id": source["id"],
        "source_name": source["name"],
        "detail": detail,
        "cells": list(cells or []),
    }


def make_quality_report(catalog: dict[str, Any], config: ImportConfig) -> dict[str, Any]:
    sources = catalog["sources"]
    source_ids = [source["id"] for source in sources]
    topics = [topic for source in sources for topic in source["topics"]]
    topic_ids = [topic["id"] for topic in topics]
    distributions = [dist for source in sources for dist in source["distributions"]]
    distribution_ids = [dist["id"] for dist in distributions]
    urls = [dist["url"] for dist in distributions]
    urls_to_sources: dict[str, set[str]] = defaultdict(set)
    warnings: list[dict[str, Any]] = []

    for source in sources:
        for dist in source["distributions"]:
            urls_to_sources[dist["url"]].add(source["id"])
            if not valid_http_url(dist["url"]):
                warnings.append(warning("malformed-url", source, dist["url"], [dist["provenance"]["cell"]]))
            target = dist.get("hyperlink_target")
            if target and target.lower().startswith(("http://", "https://")) and target != dist["url"]:
                warnings.append(
                    warning(
                        "display-hyperlink-mismatch",
                        source,
                        f"Visible URL is {dist['url']}; hyperlink target is {target}",
                        [dist["provenance"]["cell"]],
                    )
                )
        if not source["background"]:
            warnings.append(warning("missing-background", source, "Background cell is blank."))
        if not source["topics"]:
            warnings.append(warning("missing-topic", source, "No category, measures, or date fields were present."))
        if not any(topic["category"] for topic in source["topics"]):
            warnings.append(warning("missing-category", source, "No data category was supplied in the workbook."))
        if source["access_type"] == "public" and not any(
            dist["role"] == "current" for dist in source["distributions"]
        ):
            warnings.append(warning("missing-current-url", source, "Public source has no current URL."))

    shared_urls = [
        {"url": url, "source_ids": sorted(source_ids_for_url)}
        for url, source_ids_for_url in sorted(urls_to_sources.items())
        if len(source_ids_for_url) > 1
    ]
    expected_public = config.expected_public_count
    expected_request = config.expected_request_count
    checks = [
        {
            "check": "expected-public-source-count",
            "passed": catalog["summary"]["public_source_count"] == expected_public,
            "expected": expected_public,
            "actual": catalog["summary"]["public_source_count"],
        },
        {
            "check": "expected-request-only-source-count",
            "passed": catalog["summary"]["request_only_source_count"] == expected_request,
            "expected": expected_request,
            "actual": catalog["summary"]["request_only_source_count"],
        },
        {
            "check": "unique-source-ids",
            "passed": len(source_ids) == len(set(source_ids)),
            "expected": len(source_ids),
            "actual": len(set(source_ids)),
        },
        {
            "check": "unique-topic-ids",
            "passed": len(topic_ids) == len(set(topic_ids)),
            "expected": len(topic_ids),
            "actual": len(set(topic_ids)),
        },
        {
            "check": "unique-distribution-ids",
            "passed": len(distribution_ids) == len(set(distribution_ids)),
            "expected": len(distribution_ids),
            "actual": len(set(distribution_ids)),
        },
        {
            "check": "review-status-set",
            "passed": all(source["review_status"] == REVIEW_STATUS for source in sources),
            "expected": REVIEW_STATUS,
            "actual": sorted({source["review_status"] for source in sources}),
        },
        {
            "check": "source-cell-provenance",
            "passed": all(source["provenance"]["raw_cells"] for source in sources),
            "expected": "at least one raw cell per source",
            "actual": sum(bool(source["provenance"]["raw_cells"]) for source in sources),
        },
    ]
    return {
        "schema_version": SCHEMA_VERSION,
        "catalog_type": "reviewed-intake-quality-report",
        "source_workbook": catalog["source_workbook"],
        "summary": {
            **catalog["summary"],
            "current_distribution_count": sum(dist["role"] == "current" for dist in distributions),
            "archive_distribution_count": sum(dist["role"] == "archive" for dist in distributions),
            "request_info_distribution_count": sum(dist["role"] == "request_info" for dist in distributions),
            "unique_url_count": len(set(urls)),
            "warning_count": len(warnings),
            "shared_url_count": len(shared_urls),
            "missing_background_count": sum(not source["background"] for source in sources),
            "missing_category_source_count": sum(
                not any(topic["category"] for topic in source["topics"]) for source in sources
            ),
            "public_source_missing_current_url_count": sum(
                source["access_type"] == "public"
                and not any(dist["role"] == "current" for dist in source["distributions"])
                for source in sources
            ),
            "raw_cell_provenance_count": sum(
                len(source["provenance"]["raw_cells"]) for source in sources
            ),
        },
        "checks": checks,
        "passed": all(check["passed"] for check in checks),
        "duplicate_ids": {
            "sources": duplicate_values(source_ids),
            "topics": duplicate_values(topic_ids),
            "distributions": duplicate_values(distribution_ids),
        },
        "shared_urls": shared_urls,
        "warnings": warnings,
    }


def json_cell(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def source_rows(catalog: dict[str, Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for source in catalog["sources"]:
        topics = source["topics"]
        contacts = source["contacts"]
        rows.append(
            {
                "id": source["id"],
                "name": source["name"],
                "access_type": source["access_type"],
                "review_status": source["review_status"],
                "background": source["background"] or "",
                "categories_json": json_cell([topic["category"] for topic in topics if topic["category"]]),
                "measures_json": json_cell([measure for topic in topics for measure in topic["measures"]]),
                "publication_dates_json": json_cell(source["publication_dates"]),
                "data_source_dates_json": json_cell(source["data_source_dates"]),
                "current_urls_json": json_cell(
                    [dist["url"] for dist in source["distributions"] if dist["role"] == "current"]
                ),
                "archive_urls_json": json_cell(
                    [dist["url"] for dist in source["distributions"] if dist["role"] == "archive"]
                ),
                "contact_emails_json": json_cell([email for contact in contacts for email in contact["emails"]]),
                "contact_phones_json": json_cell([phone for contact in contacts for phone in contact["phones"]]),
                "contact_urls_json": json_cell([url for contact in contacts for url in contact["urls"]]),
                "notes_json": json_cell(source["notes"]),
                "source_sheet": source["provenance"]["sheet"],
                "source_rows": (
                    str(source["provenance"]["row_start"])
                    if source["provenance"]["row_start"] == source["provenance"]["row_end"]
                    else f"{source['provenance']['row_start']}-{source['provenance']['row_end']}"
                ),
            }
        )
    return rows


def topic_rows(catalog: dict[str, Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for source in catalog["sources"]:
        for topic in source["topics"]:
            rows.append(
                {
                    "id": topic["id"],
                    "source_id": source["id"],
                    "source_name": source["name"],
                    "category": topic["category"] or "",
                    "measures_raw": topic["measures_raw"] or "",
                    "measures_json": json_cell(topic["measures"]),
                    "publication_date_normalized": (
                        topic["publication_date"]["normalized"] if topic["publication_date"] else ""
                    ),
                    "publication_date_precision": (
                        topic["publication_date"]["precision"] if topic["publication_date"] else ""
                    ),
                    "data_source_date_normalized": (
                        topic["data_source_date"]["normalized"] if topic["data_source_date"] else ""
                    ),
                    "data_source_date_precision": (
                        topic["data_source_date"]["precision"] if topic["data_source_date"] else ""
                    ),
                    "source_sheet": topic["provenance"]["sheet"],
                    "source_row": topic["provenance"]["row"],
                    "source_cells_json": json_cell(topic["provenance"]["cells"]),
                }
            )
    return rows


def distribution_rows(catalog: dict[str, Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for source in catalog["sources"]:
        for dist in source["distributions"]:
            rows.append(
                {
                    "id": dist["id"],
                    "source_id": source["id"],
                    "source_name": source["name"],
                    "topic_id": dist["topic_id"] or "",
                    "role": dist["role"],
                    "format": dist["format"],
                    "url": dist["url"],
                    "display_text": dist["display_text"] or "",
                    "hyperlink_target": dist["hyperlink_target"] or "",
                    "source_sheet": dist["provenance"]["sheet"],
                    "source_row": dist["provenance"]["row"],
                    "source_cell": dist["provenance"]["cell"],
                }
            )
    return rows


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def quality_markdown(report: dict[str, Any]) -> str:
    summary = report["summary"]
    lines = [
        "# Reviewed-intake data quality report",
        "",
        f"Source workbook: `{report['source_workbook']['filename']}`  ",
        f"SHA-256: `{report['source_workbook']['sha256']}`",
        "",
        "This is an import-time review, not a claim that linked resources are currently available or authoritative.",
        "",
        "| Metric | Count |",
        "|---|---:|",
        f"| Public sources | {summary['public_source_count']} |",
        f"| Request-only sources | {summary['request_only_source_count']} |",
        f"| Total sources | {summary['source_count']} |",
        f"| Topics / measure rows | {summary['topic_count']} |",
        f"| Current links | {summary['current_distribution_count']} |",
        f"| Archived links | {summary['archive_distribution_count']} |",
        f"| Request-information links | {summary['request_info_distribution_count']} |",
        f"| Unique URLs | {summary['unique_url_count']} |",
        f"| Preserved non-empty cells | {summary['raw_cell_provenance_count']} |",
        f"| Review warnings | {summary['warning_count']} |",
        "",
        "## Validation checks",
        "",
        "| Check | Result | Expected | Actual |",
        "|---|---|---|---|",
    ]
    for check in report["checks"]:
        result = "PASS" if check["passed"] else "FAIL"
        lines.append(f"| {check['check']} | {result} | {check['expected']} | {check['actual']} |")
    lines.extend(["", "## Review warnings", ""])
    if not report["warnings"]:
        lines.append("No import-time warnings.")
    else:
        lines.append("| Issue | Source | Cells | Detail |")
        lines.append("|---|---|---|---|")
        for item in report["warnings"]:
            detail = item["detail"].replace("|", "\\|").replace("\n", " ")
            source = item["source_name"].replace("|", "\\|").replace("\n", " ")
            lines.append(
                f"| {item['issue']} | {source} | {', '.join(item['cells']) or '—'} | {detail} |"
            )
    lines.extend(
        [
            "",
            "All source records remain `imported-unverified`; URL checks and editorial verification happen in later stages.",
            "",
        ]
    )
    return "\n".join(lines)


def write_outputs(catalog: dict[str, Any], report: dict[str, Any], config: ImportConfig) -> None:
    sources_csv = source_rows(catalog)
    topics_csv = topic_rows(catalog)
    distributions_csv = distribution_rows(catalog)
    for directory in (config.catalog_dir, config.public_dir):
        write_json(directory / "reviewed-intake.json", catalog)
        write_csv(directory / "reviewed-intake.csv", sources_csv)
        write_csv(directory / "topics.csv", topics_csv)
        write_csv(directory / "distributions.csv", distributions_csv)
        write_json(directory / "data-quality-report.json", report)
    (config.catalog_dir / "data-quality-report.md").write_text(
        quality_markdown(report), encoding="utf-8"
    )


def parse_args(argv: Sequence[str] | None = None) -> ImportConfig:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path, help="Path to the SHIP .xlsx workbook")
    parser.add_argument("--catalog-dir", type=Path, default=Path("catalog"))
    parser.add_argument("--public-dir", type=Path, default=Path("public/data"))
    parser.add_argument("--expected-public-count", type=int, default=30)
    parser.add_argument("--expected-request-count", type=int, default=7)
    args = parser.parse_args(argv)
    return ImportConfig(
        input_path=args.input.resolve(),
        catalog_dir=args.catalog_dir,
        public_dir=args.public_dir,
        expected_public_count=args.expected_public_count,
        expected_request_count=args.expected_request_count,
    )


def main(argv: Sequence[str] | None = None) -> int:
    config = parse_args(argv)
    if not config.input_path.is_file():
        print(f"Input workbook not found: {config.input_path}", file=sys.stderr)
        return 2
    catalog = make_catalog(config)
    report = make_quality_report(catalog, config)
    write_outputs(catalog, report, config)
    summary = report["summary"]
    print(
        "Imported "
        f"{summary['public_source_count']} public + "
        f"{summary['request_only_source_count']} request-only sources; "
        f"{summary['topic_count']} topics; {summary['distribution_count']} distributions; "
        f"{summary['warning_count']} review warnings."
    )
    if not report["passed"]:
        failed = ", ".join(check["check"] for check in report["checks"] if not check["passed"])
        print(f"Validation failed: {failed}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
